"""
The Qumo house style, shared by every deck in this directory.

Palette, geometry and the drawing primitives. Sampled from the original
deck rather than chosen, which is why the values look arbitrary: background
#0D0E11, cards #16181C on #282B31 borders, one blue, and a monospace for
anything that labels rather than reads.

A deck imports this, calls `start(title, total)`, defines its slides, and
calls `save(path)`. Nothing in here knows what any particular deck says.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS, ART = os.path.join(HERE, "shots"), os.path.join(HERE, "art")

BG     = RGBColor(0x0D, 0x0E, 0x11)
CARD   = RGBColor(0x16, 0x18, 0x1C)
LINE   = RGBColor(0x28, 0x2B, 0x31)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
INK    = RGBColor(0xFF, 0xFF, 0xFF)
BODY   = RGBColor(0xC4, 0xCA, 0xD3)
MUTED  = RGBColor(0x8A, 0x90, 0x9B)
FOOT   = RGBColor(0x6A, 0x70, 0x80)
WARN   = RGBColor(0xFF, 0x6B, 0x7F)
PANEL  = RGBColor(0xF2, 0xF2, 0xF4)

SANS, MONO = "Arial", "Consolas"
W, H = Inches(13.333), Inches(7.5)
M = Inches(0.85)
RIGHT = Inches(12.49)
CONTENT = RIGHT - M
RULE_Y = Inches(6.90)

# Set by start(). A deck is a sequence of side effects on one presentation,
# so this is module state rather than an object nobody would subclass.
prs = None
BLANK = None
_page = {"n": 0}
TOTAL = 0
DECK = ""


def start(deck_title, total):
    """Open a presentation. `deck_title` sits beside the mark on every slide."""
    global prs, BLANK, TOTAL, DECK
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    BLANK = prs.slide_layouts[6]
    _page["n"] = 0
    TOTAL = total
    DECK = deck_title
    return prs


def save(path):
    prs.save(path)
    return _page["n"]


# ---------------- primitives ----------------

def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(1, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background(); bg.shadow.inherit = False
    return s


def text(s, txt, x, y, w, h, size=13, bold=False, colour=BODY, font=SANS,
         align=PP_ALIGN.LEFT, spacing=1.5, spc=None, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(txt.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold
        r.font.name = font; r.font.color.rgb = colour
        if spc:
            # Letter-spacing, which python-pptx does not expose. The tracked
            # out mono labels look wrong without it.
            r.font._rPr.set("spc", str(int(spc * 100)))
    return tb


def box(s, x, y, w, h, fill=CARD, border=LINE, radius=0.045):
    sh = s.shapes.add_shape(5, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = border; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    sh.adjustments[0] = radius
    return sh


def hline(s, x, y, w, colour=LINE, thick=Emu(9525)):
    sh = s.shapes.add_shape(1, x, y, w, thick)
    sh.fill.solid(); sh.fill.fore_color.rgb = colour
    sh.line.fill.background(); sh.shadow.inherit = False


def logo(s, x, y, h, variant="white"):
    path = os.path.join(ART, f"qumo-{variant}.png")
    iw, ih = Image.open(path).size
    s.shapes.add_picture(path, x, y, height=h, width=Emu(int(h * iw / ih)))


def masthead(s):
    logo(s, M, Inches(0.40), Inches(0.21))
    text(s, DECK, M + Inches(0.68), Inches(0.41), Inches(4), Inches(0.24),
         size=10.5, colour=MUTED, spacing=1.0)


def footer(s):
    _page["n"] += 1
    hline(s, M, RULE_Y, CONTENT)
    text(s, "Qumo  ·  indicative, for discussion", M, RULE_Y + Inches(0.12),
         Inches(8), Inches(0.24), size=8, colour=FOOT, font=MONO, spacing=1.0)
    text(s, f"{_page['n']:02d} / {TOTAL}", RIGHT - Inches(2), RULE_Y + Inches(0.12),
         Inches(2), Inches(0.24), size=8, colour=FOOT, font=MONO,
         align=PP_ALIGN.RIGHT, spacing=1.0)


def kicker(s, label, colour=ACCENT, y=Inches(1.00)):
    hline(s, M, y + Inches(0.055), Inches(0.30), colour, Emu(19050))
    text(s, label.upper(), M + Inches(0.41), y - Inches(0.055), Inches(9), Inches(0.26),
         size=9, colour=colour, font=MONO, spacing=1.0, spc=1.6)


def est_lines(txt, size_pt, width_in):
    """
    How many lines a string will wrap to, near enough.

    Here because a heading placed at a fixed height is correct for one line
    and wrong for two: an early draft shipped nine slides with the second
    line of the heading printed through the paragraph under it. Arial Bold
    advances about 0.52 of the point size per character.

    A newline in the string is a line the author asked for, and each side
    of it wraps on its own. Counting the whole string as one stream missed
    that, which is how a three-line heading was placed as a two-line one.
    """
    per_char = 0.52 * size_pt / 72
    per_line = max(1, int(width_in / per_char))
    lines = 0
    for segment in txt.split("\n"):
        lines += 1
        cur = 0
        for word in segment.split():
            add = len(word) + (1 if cur else 0)
            if cur + add > per_line:
                lines += 1; cur = len(word)
            else:
                cur += add
    return lines


# Arial's line box is about 1.15 times the point size before any line
# spacing multiple is applied, so a heading set at 1.2 spacing occupies
# nearer 1.38 of its size per line. Reserving only 1.2 left the paragraph
# under a wrapped heading sitting on its last line.
LINE_BOX = 1.15


def h1(s, txt, y=Inches(1.46), size=29, w=None, colour=INK):
    lines = est_lines(txt, size, (w or CONTENT) / Inches(1))
    h = Inches(lines * size * 1.2 * LINE_BOX / 72)
    text(s, txt, M, y, w or CONTENT, h, size=size, bold=True, colour=colour, spacing=1.2)
    return y + h


def shot(s, name, x, y, max_w, max_h, panel=True, fallback=None):
    """
    A screenshot on a light panel, which is how light UI reads on black.

    `fallback` is for screens that only exist when capture.mjs was given the
    optional arguments: the deck still builds without them rather than
    dying on a missing file halfway through.
    """
    path = os.path.join(SHOTS, f"{name}.png")
    if not os.path.exists(path) and fallback:
        path = os.path.join(SHOTS, f"{fallback}.png")
    iw, ih = Image.open(path).size
    pad = Inches(0.09) if panel else Inches(0)
    scale = min((max_w - 2 * pad) / iw, (max_h - 2 * pad) / ih)
    w, h = Emu(int(iw * scale)), Emu(int(ih * scale))
    cx = x + Emu(int((max_w - w - 2 * pad) / 2)) + pad
    if panel:
        box(s, cx - pad, y, w + 2 * pad, h + 2 * pad, fill=PANEL, border=None, radius=0.03)
    s.shapes.add_picture(path, cx, y + pad, width=w, height=h)
    return h + 2 * pad


def stats(s, items, y, x=None, total_w=None, gap=Inches(0.16), hgt=Inches(0.95)):
    x = x if x is not None else M
    total_w = total_w if total_w is not None else CONTENT
    n = len(items)
    w = Emu(int((total_w - gap * (n - 1)) / n))
    for i, (value, label) in enumerate(items):
        cx = x + Emu(int((w + gap) * i))
        box(s, cx, y, w, hgt)
        text(s, value, cx + Inches(0.22), y + Inches(0.15), w - Inches(0.44), Inches(0.42),
             size=22, bold=True, colour=ACCENT if i == 0 else INK, spacing=1.0)
        text(s, label.upper(), cx + Inches(0.22), y + Inches(0.62), w - Inches(0.44), Inches(0.24),
             size=8, colour=MUTED, font=MONO, spacing=1.0, spc=1.2)


def dotted(s, items, x, y, w, size=11.5, gap=Inches(0.255), colour=BODY):
    for i, item in enumerate(items):
        top = y + Emu(int(gap * i))
        text(s, "·", x, top, Inches(0.16), gap, size=size, colour=MUTED, spacing=1.35)
        text(s, item, x + Inches(0.21), top, w - Inches(0.21), gap, size=size,
             colour=colour, spacing=1.35)
    return y + Emu(int(gap * len(items)))


def rows(s, pairs, y, label_w=Inches(3.3), gap=Inches(0.84), size=12):
    for i, (label, value) in enumerate(pairs):
        top = y + Emu(int(gap * i))
        hline(s, M, top, CONTENT)
        text(s, label, M, top + Inches(0.14), label_w, Inches(0.3), size=size,
             bold=True, colour=INK, spacing=1.35)
        text(s, value, M + label_w + Inches(0.1), top + Inches(0.14),
             CONTENT - label_w - Inches(0.1), gap - Inches(0.2), size=size,
             colour=BODY, spacing=1.45)


def tri_cards(s, cards, y, hgt=Inches(3.45), highlight=0):
    gap = Inches(0.22)
    w = Emu(int((CONTENT - gap * 2) / 3))
    for i, c in enumerate(cards):
        x = M + Emu(int((w + gap) * i))
        on = i == highlight
        box(s, x, y, w, hgt, border=ACCENT if on else LINE)
        px, pw = x + Inches(0.26), w - Inches(0.52)
        text(s, c["tag"].upper(), px, y + Inches(0.24), pw, Inches(0.24), size=8,
             colour=ACCENT if on else MUTED, font=MONO, spacing=1.0, spc=1.2)
        text(s, c["title"], px, y + Inches(0.54), pw, Inches(0.42), size=19, bold=True,
             colour=INK, spacing=1.1)
        text(s, c["sub"], px, y + Inches(1.02), pw, Inches(0.28), size=11, colour=BODY, spacing=1.2)
        points_y, points_gap = y + Inches(1.32), Inches(0.225)
        divider = y + hgt - Inches(0.92)
        bottom = points_y + Emu(int(points_gap * len(c["points"])))
        if bottom > divider:
            # The footer block hangs off the bottom edge so all three cards
            # line their prices up, which means a card that is too short for
            # its own bullets prints them through the divider. That defect is
            # invisible in the code and obvious on the page, so fail here.
            short = (bottom - divider) / Inches(1)
            raise ValueError(
                f"tri_cards: card {i} ({c['title']}) needs about "
                f"{short + 0.05:.2f}in more height for {len(c['points'])} points"
            )
        dotted(s, c["points"], px, points_y, pw, size=10, gap=points_gap)
        hline(s, px, divider, pw)
        text(s, c["foot_label"].upper(), px, y + hgt - Inches(0.80), pw, Inches(0.24),
             size=8, colour=ACCENT if on else MUTED, font=MONO, spacing=1.0, spc=1.2)
        text(s, c["foot"], px, y + hgt - Inches(0.54), pw, Inches(0.3), size=14, bold=True,
             colour=ACCENT if on and c.get("foot_accent") else INK, spacing=1.15)
        if c.get("foot_note"):
            text(s, c["foot_note"], px, y + hgt - Inches(0.26), pw, Inches(0.24), size=9,
                 colour=MUTED, spacing=1.1)


def two_col(s, kick, title, paras, shot_name, shot_y=Inches(1.75), shot_h=Inches(4.3)):
    masthead(s); kicker(s, kick)
    y = h1(s, title, w=Inches(4.9))
    ty = y + Inches(0.22)
    for i, p in enumerate(paras):
        size = 13 if i == 0 else 12
        hh = Inches(est_lines(p, size, 4.6) * size * 1.55 / 72)
        text(s, p, M, ty, Inches(4.75), hh, size=size,
             colour=BODY if i == 0 else MUTED, spacing=1.55)
        ty += hh + Inches(0.26)
    shot(s, shot_name, Inches(6.06), shot_y, Inches(6.43), shot_h)
    footer(s)


def phones(s, items, y=Inches(1.70), h=Inches(3.85), x0=Inches(6.06), total=Inches(6.43)):
    gap = Inches(0.14)
    w = Emu(int((total - gap * (len(items) - 1)) / len(items)))
    for i, (name, cap, sub) in enumerate(items):
        x = x0 + Emu(int((w + gap) * i))
        used = shot(s, name, x, y, w, h, fallback="shopper-wallet")
        text(s, cap, x, y + used + Inches(0.16), w, Inches(0.26), size=12, bold=True,
             colour=INK, spacing=1.2)
        text(s, sub, x, y + used + Inches(0.44), w, Inches(0.26), size=10.5,
             colour=MUTED, spacing=1.2)


def part(kick, title, body):
    s = slide(); masthead(s)
    kicker(s, kick, y=Inches(2.56))
    h1(s, title, y=Inches(3.02), size=44)
    text(s, body, M, Inches(4.62), Inches(9.6), Inches(0.9), size=13, colour=BODY, spacing=1.55)
    footer(s)
